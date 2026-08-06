"""Safe deterministic manuscript analysis, author review, and artifact generation."""

from __future__ import annotations

import hashlib
import json
import re
import sqlite3
import subprocess
import tempfile
import textwrap
import uuid
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Iterable
from contextlib import closing
from datetime import UTC, datetime
from html import escape, unescape
from io import BytesIO
from pathlib import Path
from typing import Any, Literal, Protocol

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from fastapi import HTTPException
from matplotlib.mathtext import math_to_image
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from journal_matcher_api.foundation import FoundationStore, Principal

Decision = Literal["accepted", "rejected", "modified"]
CATEGORY_COLORS = {
    "scope-fit": "2F5597",
    "literature-positioning": "7030A0",
    "novelty-significance": "C00000",
    "form": "1F4E79",
    "language": "1F4E79",
    "structure": "7030A0",
    "content": "2F5597",
    "scientific-question": "C00000",
    "compliance": "C65911",
    "journal-format": "C65911",
    "methodology-reporting": "548235",
    "scientific-concern": "C00000",
    "unresolved": "666666",
    "scientific-framing": "2F5597",
    "theory-methodology": "548235",
    "validation-robustness": "C00000",
    "results-analysis": "2F5597",
    "figures-equations": "7030A0",
    "writing": "1F4E79",
}


class QualitativeModelAdapter(Protocol):
    """Provider-neutral boundary for a future schema-constrained qualitative model."""

    def analyze(
        self, manuscript_segments: list[dict[str, str]], profile_claims: list[dict[str, object]]
    ) -> list[dict[str, object]]:
        """Return untrusted candidate recommendations for deterministic validation."""


def now_iso() -> str:
    return datetime.now(UTC).isoformat()


def stable_id(*parts: object) -> str:
    digest = hashlib.sha256(json.dumps(parts, sort_keys=True, default=str).encode()).hexdigest()
    return str(uuid.UUID(digest[:32]))


def scope_analysis_records(
    analysis_id: str, record_type: str, records: list[dict[str, object]]
) -> list[dict[str, object]]:
    """Give child records an analysis-scoped identity for safe cross-project reuse."""
    scoped: list[dict[str, object]] = []
    for record in records:
        original_id = record.get("id")
        if not isinstance(original_id, str) or not original_id:
            raise HTTPException(status_code=422, detail=f"{record_type} record is missing an identity")
        item = dict(record)
        item["id"] = stable_id(analysis_id, record_type, original_id)
        scoped.append(item)
    return scoped


def scientific_invariants(text: str) -> dict[str, list[str]]:
    return {
        "numbers": re.findall(r"(?<!\w)[+-]?\d+(?:[.,]\d+)?%?", text),
        "units": re.findall(r"\b\d+(?:[.,]\d+)?\s*(?:mg|g|kg|mL|L|mm|cm|m|°C|K|Pa|kPa|MPa|%)\b", text),
        "citations": re.findall(r"\[[0-9,; -]+\]|\([A-Z][A-Za-z-]+(?: et al\.)?,? \d{4}[a-z]?\)", text),
        "equations": re.findall(r"[^\n]{0,30}[=≈≤≥][^\n]{0,30}", text),
    }


def extract_official_rules(snapshots: list[dict[str, str]]) -> list[dict[str, object]]:
    rules: list[dict[str, object]] = []
    for snapshot in snapshots:
        content = snapshot["content"]
        source_id = snapshot["source_id"]
        abstract_limit_pattern = r"abstract.{0,100}?(?:maximum|limit|not exceed|up to)?\s*(\d{2,4})\s*words"
        for match in re.finditer(abstract_limit_pattern, content, re.I):
            rules.append(
                {
                    "id": stable_id(source_id, "abstract-word-limit", match.group(1)),
                    "key": "abstract-word-limit",
                    "value": int(match.group(1)),
                    "sourceId": source_id,
                    "locator": f"characters:{match.start()}-{match.end()}",
                    "status": "validated",
                }
            )
        article_limit_pattern = r"letters?\s*\(\s*length limit\s*:\s*(\d{3,5})\s*words"
        for match in re.finditer(article_limit_pattern, content, re.I):
            rules.append(
                {
                    "id": stable_id(source_id, "article-word-limit", match.group(1)),
                    "key": "article-word-limit",
                    "value": int(match.group(1)),
                    "sourceId": source_id,
                    "locator": f"characters:{match.start()}-{match.end()}",
                    "status": "validated",
                }
            )
        end_matter_pattern = r"up to\s*(\w+)\s*pages?\s+of\s+(?:appendices|end matter)"
        for match in re.finditer(end_matter_pattern, content, re.I):
            value = {"one": 1, "two": 2, "three": 3}.get(match.group(1).casefold())
            if value:
                rules.append(
                    {
                        "id": stable_id(source_id, "end-matter-page-limit", value),
                        "key": "end-matter-page-limit",
                        "value": value,
                        "sourceId": source_id,
                        "locator": f"characters:{match.start()}-{match.end()}",
                        "status": "validated",
                    }
                )
        requirements = {
            "ethics-statement": ("ethics statement", "ethical approval"),
            "conflict-statement": ("conflict of interest", "competing interests"),
            "funding-statement": ("funding statement", "funding information"),
            "data-availability": ("data availability", "availability of data"),
        }
        folded = content.casefold()
        for key, phrases in requirements.items():
            found = next((phrase for phrase in phrases if phrase in folded), None)
            if found:
                start = folded.index(found)
                rules.append(
                    {
                        "id": stable_id(source_id, key),
                        "key": key,
                        "value": True,
                        "sourceId": source_id,
                        "locator": f"characters:{start}-{start + len(found)}",
                        "status": "validated",
                    }
                )
    unique: dict[str, dict[str, object]] = {}
    for rule in rules:
        key = str(rule["key"])
        if key in unique and unique[key]["value"] != rule["value"]:
            unique[key]["status"] = "conflict"
            rule["status"] = "conflict"
        unique.setdefault(key, rule)
    return list(unique.values())


def build_recommendations(
    manuscript_text: str,
    profile: dict[str, object],
    rules: list[dict[str, object]],
    anchor: str,
    manuscript_segments: list[dict[str, str]] | None = None,
    journal_title: str = "the target journal",
) -> list[dict[str, object]]:
    folded = manuscript_text.casefold()
    prompt_like_upload = any(
        marker in folded for marker in ("ignore all prior", "ignore previous instructions", "system prompt")
    )
    recommendations: list[dict[str, object]] = []

    def add(
        key: str,
        category: str,
        severity: str,
        rationale: str,
        basis: str,
        source_ids: list[str],
        proposed: str | None,
        scientific_impact: bool = False,
        original: str = "",
        recommendation_anchor: str | None = None,
        confidence: float | None = None,
    ) -> None:
        target_anchor = recommendation_anchor or anchor
        recommendations.append(
            {
                "id": stable_id(profile["id"], key, target_anchor),
                "key": key,
                "anchor": target_anchor,
                "originalText": original,
                "proposedText": proposed,
                "category": category,
                "severity": severity,
                "rationale": rationale,
                "basis": basis,
                "sourceIds": source_ids,
                "evidenceCoverage": f"{len(source_ids)} source(s)",
                "confidence": (
                    confidence if confidence is not None else 1.0 if basis == "official-requirement" else 0.75
                ),
                "uncertainty": "Heuristic or deterministic review; the author must verify scientific meaning.",
                "scientificImpact": scientific_impact,
                "decision": "pending",
            }
        )

    for rule in rules:
        if rule["status"] != "validated":
            continue
        source_ids = [str(rule["sourceId"])]
        key = str(rule["key"])
        if key == "abstract-word-limit":
            abstract = _section_text(manuscript_text, "abstract")
            limit = int(str(rule["value"]))
            if abstract and len(abstract.split()) > limit:
                add(
                    key,
                    "journal-format",
                    "required",
                    f"The abstract has {len(abstract.split())} words; the validated official limit is {limit}.",
                    "official-requirement",
                    source_ids,
                    f"Shorten the abstract to at most {limit} words without changing results or numeric values.",
                    True,
                )
        elif key == "article-word-limit":
            limit = int(str(rule["value"]))
            core_text = re.split(r"\n\s*(?:references|bibliography)\b", manuscript_text, maxsplit=1, flags=re.I)[0]
            word_count = len(re.findall(r"\b[A-Za-z][A-Za-z'-]*\b", core_text))
            if word_count > limit:
                add(
                    key,
                    "journal-format",
                    "required",
                    f"The extracted core contains approximately {word_count} words; "
                    f"{journal_title} limits this manuscript type's core to {limit} words.",
                    "official-requirement",
                    source_ids,
                    "Reduce the core to the official limit; move specialist derivations to End Matter or "
                    "Supplemental Material and verify with the APS length tool.",
                    True,
                    original=f"Approximate extracted core word count: {word_count}",
                )
        elif key == "end-matter-page-limit":
            continue
        else:
            labels = {
                "ethics-statement": ("ethics", "Add or identify the ethics approval/exemption statement."),
                "conflict-statement": ("conflict", "Add or identify the competing-interests statement."),
                "funding-statement": ("funding", "Add or identify the funding statement."),
                "data-availability": ("data availability", "Add or identify the data-availability statement."),
            }
            needle, action = labels[key]
            if needle not in folded:
                add(key, "journal-format", "required", action, "official-requirement", source_ids, None, True)

    claims = profile.get("claims", [])
    metric_claims: dict[str, dict[str, object]] = {}
    if isinstance(claims, list):
        for claim in claims:
            if not isinstance(claim, dict):
                continue
            claim_key = str(claim.get("key", ""))
            if claim_key.startswith("metric:"):
                metric_claims[claim_key] = claim
                continue
            if not claim_key.startswith("section:"):
                continue
            section = str(claim["key"]).split(":", 1)[1]
            coverage = str(claim.get("coverage", ""))
            numbers = re.search(r"(\d+)\s+of\s+(\d+)", coverage)
            support_ratio = int(numbers.group(1)) / int(numbers.group(2)) if numbers else 0.0
            if support_ratio >= 0.5 and re.search(rf"\b{re.escape(section)}\b", folded) is None:
                add(
                    f"observed-section:{section}",
                    "structure",
                    "strongly-recommended",
                    f"A labeled {section} section appeared in {claim.get('coverage', 'the reference sample')}.",
                    "observed-pattern",
                    [str(item) for item in claim.get("sourceIds", [])],
                    f"Consider whether a distinct {section.title()} section would improve alignment.",
                    confidence=round(support_ratio, 2),
                )

    manuscript_words = re.findall(r"\b[A-Za-z][A-Za-z'-]*\b", manuscript_text)
    manuscript_sentences = [part for part in re.split(r"[.!?]+", manuscript_text) if len(part.split()) >= 3]
    manuscript_mean_sentence = (
        sum(len(item.split()) for item in manuscript_sentences) / len(manuscript_sentences)
        if manuscript_sentences
        else 0.0
    )
    manuscript_equation_density = (
        len(re.findall(r"\(\d{1,3}\)", manuscript_text)) * 1000 / max(len(manuscript_words), 1)
    )
    for metric_key, actual, category, action in (
        (
            "metric:meanSentenceWords",
            manuscript_mean_sentence,
            "language",
            "Shorten syntactically dense sentences while preserving equations, qualifications, and causal logic.",
        ),
        (
            "metric:equationMarkersPerThousandWords",
            manuscript_equation_density,
            "structure",
            "Reduce main-text derivation density and move intermediate steps to specialist material.",
        ),
    ):
        claim = metric_claims.get(metric_key)
        value = claim.get("value") if claim else None
        if not isinstance(value, dict):
            continue
        observed_max = float(value.get("max", 0.0))
        if observed_max > 0 and actual > observed_max * 1.2:
            claim_source_ids = claim.get("sourceIds", []) if isinstance(claim, dict) else []
            add(
                f"deviation:{metric_key.split(':', 1)[1]}",
                category,
                "strongly-recommended",
                f"The manuscript value ({actual:.1f}) exceeds the observed sample maximum ({observed_max:.1f}).",
                "observed-pattern",
                [str(item) for item in claim_source_ids] if isinstance(claim_source_ids, list) else [],
                action,
                confidence=0.8,
            )

    title = next((line.strip() for line in manuscript_text.splitlines() if len(line.strip()) > 8), "")
    abstract = _section_text(manuscript_text, "abstract") or "\n".join(manuscript_text.splitlines()[:35])
    introduction = _section_text(manuscript_text, "introduction")
    background = _section_text(manuscript_text, "background")
    conclusion = _section_text(manuscript_text, "conclusion")

    def located(needle: str, fallback: str) -> tuple[str, str]:
        if prompt_like_upload:
            return fallback, ""
        for segment in manuscript_segments or []:
            segment_text = str(segment.get("text", ""))
            if needle.casefold() in segment_text.casefold():
                position = segment_text.casefold().find(needle.casefold())
                excerpt = segment_text[max(0, position - 100) : position + 300].strip()
                return str(segment.get("anchor", fallback)), excerpt
        return fallback, ""

    title_anchor, title_excerpt = located(title, anchor)
    if "coherence" in abstract.casefold() and "coherence" not in title.casefold():
        add(
            "title-central-result",
            "language",
            "strongly-recommended",
            f"{journal_title} expects the title to convey the most important result; the current title may not "
            "foreground the central result.",
            "expert-suggestion",
            [],
            "Rewrite the title to foreground the demonstrated decomposition of coherence into heat and work; "
            "preserve the precise scope and avoid overstating generality.",
            True,
            original=title_excerpt or title,
            recommendation_anchor=title_anchor,
        )
    abstract_words = len(re.findall(r"\b[A-Za-z][A-Za-z'-]*\b", abstract))
    if abstract_words > 180:
        abstract_anchor, abstract_excerpt = located("In quantum thermodynamics", anchor)
        add(
            "abstract-concision",
            "language",
            "strongly-recommended",
            f"The abstract contains approximately {abstract_words} words; {journal_title} expects a concise "
            "statement of the principal result for its readership.",
            "expert-suggestion",
            [],
            "Compress background and repeated interpretation while retaining the problem, method, principal "
            "result, validity order, and significance.",
            True,
            original=abstract_excerpt,
            recommendation_anchor=abstract_anchor,
        )
    intro_words = len(re.findall(r"\b[A-Za-z][A-Za-z'-]*\b", introduction))
    if intro_words > 900:
        intro_anchor, intro_excerpt = located("INTRODUCTION", anchor)
        add(
            "introduction-broad-reader",
            "structure",
            "strongly-recommended",
            f"The extracted Introduction is approximately {intro_words} words, which may compete with the "
            f"space available for the central {journal_title} result.",
            "expert-suggestion",
            [],
            "Condense the field survey, identify the unresolved contradiction earlier, and state the paper's "
            "decisive advance in the opening narrative.",
            True,
            original=intro_excerpt,
            recommendation_anchor=intro_anchor,
        )
    background_words = len(re.findall(r"\b[A-Za-z][A-Za-z'-]*\b", background))
    if background_words > 700:
        background_anchor, background_excerpt = located("BACKGROUND", anchor)
        add(
            "background-to-end-matter",
            "structure",
            "strongly-recommended",
            f"The standalone Background section is approximately {background_words} words and contains "
            "specialist derivations.",
            "expert-suggestion",
            [],
            "Keep only definitions essential to establish the result in the Letter; move extended framework "
            "comparisons and derivations to End Matter or Supplemental Material.",
            True,
            original=background_excerpt,
            recommendation_anchor=background_anchor,
        )
    equation_count = len(re.findall(r"\(\d{1,3}\)", manuscript_text))
    if equation_count > 30:
        equations_anchor, equations_excerpt = located("(1)", anchor)
        add(
            "derivation-density",
            "structure",
            "strongly-recommended",
            f"At least {equation_count} numbered-equation references were detected, indicating derivation "
            "density inconsistent with a concise Letter core.",
            "expert-suggestion",
            [],
            "Retain the equations needed to establish the physical claim and relocate intermediate algebra "
            "to the permitted specialist material.",
            True,
            original=equations_excerpt,
            recommendation_anchor=equations_anchor,
        )
    if "second order" in folded and not re.search(r"\b(convergence|error bound|range of validity|benchmark)\b", folded):
        validity_anchor, validity_excerpt = located("second order", anchor)
        add(
            "perturbative-validity",
            "scientific-concern",
            "question",
            "The central conclusion is based on truncation at second order, but the deterministic review did "
            "not locate a convergence test, error bound, benchmark, or explicit range-of-validity analysis.",
            "expert-suggestion",
            [],
            "Demonstrate or delimit the regime in which second order is sufficient, ideally with a solvable "
            "benchmark, numerical comparison, or quantitative remainder estimate.",
            True,
            original=validity_excerpt,
            recommendation_anchor=validity_anchor,
        )
    if not re.search(r"\b(limitations?|restricted to|valid(?:ity)? regime|breaks down)\b", folded):
        add(
            "scientific-limitations",
            "scientific-concern",
            "question",
            f"Readers of {journal_title} need to know how general the central claim is.",
            "expert-suggestion",
            [],
            "State the assumptions and boundaries explicitly: coupling/dynamical regime, differentiability, "
            "spectral conditions, perturbative parameter, initial states, and cases not covered.",
            True,
        )
    if conclusion and not re.search(r"\b(future|outlook|open question|remains)\b", conclusion, re.I):
        conclusion_anchor, conclusion_excerpt = located("CONCLUSION", anchor)
        add(
            "conclusion-outlook",
            "structure",
            "optional",
            f"The official {journal_title} guidance asks the conclusion to summarize results and point to future directions.",
            "expert-suggestion",
            [],
            "Add one restrained outlook sentence identifying the most consequential test or extension, "
            "without introducing unsupported claims.",
            True,
            original=conclusion_excerpt,
            recommendation_anchor=conclusion_anchor,
        )
    add(
        "scope-fit-author-review",
        "scope-fit",
        "question",
        "Scope fit requires scientific judgment and cannot be established from keyword overlap alone.",
        "expert-suggestion",
        [],
        f"Explain which {journal_title} acceptance and scope expectations are met and why the result matters "
        "beyond the immediate specialty.",
        True,
        original="" if prompt_like_upload else abstract[:400],
    )
    bibliography = re.split(r"\n\s*(?:references|bibliography)\b", manuscript_text, maxsplit=1, flags=re.I)
    bibliography_text = bibliography[1] if len(bibliography) == 2 else ""
    reference_entries = re.findall(r"(?m)^\s*(?:\[\d+\]|\d+[.)])\s+", bibliography_text)
    add(
        "literature-positioning-review",
        "literature-positioning",
        "question",
        (
            f"The deterministic extraction found approximately {len(reference_entries)} numbered bibliography "
            "entries; topical coverage, close prior work, and novelty still require source-by-source review."
            if bibliography_text
            else "No reliably delimited References or Bibliography section was found in the extracted manuscript."
        ),
        "expert-suggestion",
        [],
        "Verify every central claim against the cited bibliography and recent literature; identify the closest "
        "work explicitly, state the non-overlapping advance, add missing citations, and narrow any priority claim "
        "that cannot be supported.",
        True,
        original=bibliography_text[:500],
    )
    add(
        "novelty-significance-gate",
        "novelty-significance",
        "question",
        "A journal-level novelty claim must distinguish a new result from a new presentation of known equations, "
        "known limiting behavior, or a direct corollary of cited work.",
        "expert-suggestion",
        [],
        "Write a one-sentence novelty claim naming the closest result, the exact technical difference, the new "
        "evidence supplied here, and the consequence for a broader physics audience. List claims that should be "
        "avoided because the literature already supports them.",
        True,
        original="" if prompt_like_upload else abstract[:500],
    )
    return recommendations


def _section_text(text: str, section: str) -> str:
    match = re.search(rf"\b{section}\b\s*[:\n]?(.+?)(?=\n\s*[A-Z][A-Za-z ]{{2,30}}\s*[:\n]|$)", text, re.I | re.S)
    return match.group(1).strip() if match else ""


class AnalysisRepository:
    def __init__(self, database_path: Path) -> None:
        self.database_path = database_path

    def migrate(self) -> None:
        with closing(sqlite3.connect(self.database_path)) as connection, connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS analysis_runs (
                    id TEXT PRIMARY KEY, project_id TEXT NOT NULL, workspace_id TEXT NOT NULL,
                    profile_version_id TEXT NOT NULL, manuscript_hash TEXT NOT NULL,
                    invariant_json TEXT NOT NULL, limitations_json TEXT NOT NULL,
                    status TEXT NOT NULL, created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS guide_rules (
                    id TEXT PRIMARY KEY, analysis_id TEXT NOT NULL, rule_json TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS recommendations (
                    id TEXT PRIMARY KEY, analysis_id TEXT NOT NULL, recommendation_json TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS recommendation_decisions (
                    id TEXT PRIMARY KEY, recommendation_id TEXT NOT NULL, workspace_id TEXT NOT NULL,
                    decision TEXT NOT NULL, modified_text TEXT, created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS analysis_artifacts (
                    id TEXT PRIMARY KEY, analysis_id TEXT NOT NULL, workspace_id TEXT NOT NULL,
                    kind TEXT NOT NULL, object_key TEXT NOT NULL, content_hash TEXT NOT NULL,
                    validation_json TEXT NOT NULL, created_at TEXT NOT NULL,
                    UNIQUE(analysis_id, kind)
                );
                """
            )

    def create(
        self,
        principal: Principal,
        project_id: str,
        profile_version_id: str,
        manuscript_hash: str,
        invariants: dict[str, list[str]],
        rules: list[dict[str, object]],
        recommendations: list[dict[str, object]],
        limitations: list[str],
    ) -> dict[str, object]:
        analysis_id = stable_id(principal.workspace_id, project_id, profile_version_id, manuscript_hash)
        scoped_rules = scope_analysis_records(analysis_id, "guide-rule", rules)
        scoped_recommendations = scope_analysis_records(analysis_id, "recommendation", recommendations)
        with closing(sqlite3.connect(self.database_path)) as connection, connection:
            existing = connection.execute("SELECT id FROM analysis_runs WHERE id = ?", (analysis_id,)).fetchone()
            if not existing:
                connection.execute(
                    "INSERT INTO analysis_runs VALUES (?, ?, ?, ?, ?, ?, ?, 'review', ?)",
                    (
                        analysis_id,
                        project_id,
                        principal.workspace_id,
                        profile_version_id,
                        manuscript_hash,
                        json.dumps(invariants, sort_keys=True),
                        json.dumps(limitations),
                        now_iso(),
                    ),
                )
                connection.executemany(
                    "INSERT INTO guide_rules VALUES (?, ?, ?)",
                    [(str(rule["id"]), analysis_id, json.dumps(rule, sort_keys=True)) for rule in scoped_rules],
                )
                connection.executemany(
                    "INSERT INTO recommendations VALUES (?, ?, ?)",
                    [
                        (str(item["id"]), analysis_id, json.dumps(item, sort_keys=True))
                        for item in scoped_recommendations
                    ],
                )
        return self.get(principal, analysis_id)

    def get(self, principal: Principal, analysis_id: str) -> dict[str, object]:
        with closing(sqlite3.connect(self.database_path)) as connection:
            connection.row_factory = sqlite3.Row
            run = connection.execute(
                "SELECT * FROM analysis_runs WHERE id = ? AND workspace_id = ?", (analysis_id, principal.workspace_id)
            ).fetchone()
            if run is None:
                raise HTTPException(status_code=404, detail="Analysis not found")
            rules = [
                json.loads(row[0])
                for row in connection.execute("SELECT rule_json FROM guide_rules WHERE analysis_id = ?", (analysis_id,))
            ]
            recommendations = []
            for row in connection.execute(
                "SELECT id, recommendation_json FROM recommendations WHERE analysis_id = ?", (analysis_id,)
            ):
                item = json.loads(row["recommendation_json"])
                decision = connection.execute(
                    """SELECT decision, modified_text, created_at FROM recommendation_decisions
                       WHERE recommendation_id = ? AND workspace_id = ? ORDER BY created_at DESC LIMIT 1""",
                    (row["id"], principal.workspace_id),
                ).fetchone()
                if decision:
                    item["decision"] = decision["decision"]
                    item["modifiedText"] = decision["modified_text"]
                    item["decidedAt"] = decision["created_at"]
                recommendations.append(item)
            artifacts = [
                dict(row)
                for row in connection.execute(
                    "SELECT id, kind, content_hash, validation_json, created_at FROM analysis_artifacts "
                    "WHERE analysis_id = ? AND workspace_id = ? ORDER BY kind",
                    (analysis_id, principal.workspace_id),
                )
            ]
        return {
            "id": run["id"],
            "projectId": run["project_id"],
            "profileVersionId": run["profile_version_id"],
            "status": run["status"],
            "limitations": json.loads(run["limitations_json"]),
            "rules": rules,
            "recommendations": recommendations,
            "artifacts": artifacts,
            "createdAt": run["created_at"],
        }

    def latest_for_project(self, principal: Principal, project_id: str) -> dict[str, object]:
        with closing(sqlite3.connect(self.database_path)) as connection:
            row = connection.execute(
                "SELECT id FROM analysis_runs WHERE project_id = ? AND workspace_id = ? "
                "ORDER BY created_at DESC LIMIT 1",
                (project_id, principal.workspace_id),
            ).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Analysis not found")
        return self.get(principal, str(row[0]))

    def add_ai_review(
        self,
        principal: Principal,
        analysis_id: str,
        recommendations: list[dict[str, object]],
        limitations: list[str],
    ) -> dict[str, object]:
        """Append an idempotent AI review to an existing workspace-scoped analysis."""
        analysis = self.get(principal, analysis_id)
        merged_limitations = list(dict.fromkeys([*analysis["limitations"], *limitations]))  # type: ignore[misc]
        scoped_recommendations = scope_analysis_records(analysis_id, "recommendation", recommendations)
        with closing(sqlite3.connect(self.database_path)) as connection, connection:
            connection.executemany(
                "INSERT OR IGNORE INTO recommendations VALUES (?, ?, ?)",
                [(str(item["id"]), analysis_id, json.dumps(item, sort_keys=True)) for item in scoped_recommendations],
            )
            connection.execute(
                "UPDATE analysis_runs SET limitations_json = ? WHERE id = ? AND workspace_id = ?",
                (json.dumps(merged_limitations), analysis_id, principal.workspace_id),
            )
        return self.get(principal, analysis_id)

    def decide(
        self,
        principal: Principal,
        analysis_id: str,
        recommendation_id: str,
        decision: Decision,
        modified_text: str | None,
    ) -> dict[str, object]:
        analysis = self.get(principal, analysis_id)
        recommendation_items = analysis["recommendations"]
        if not isinstance(recommendation_items, list):
            raise HTTPException(status_code=500, detail="Stored recommendation set is invalid")
        recommendation = next(
            (item for item in recommendation_items if isinstance(item, dict) and item["id"] == recommendation_id),
            None,
        )
        if recommendation is None:
            raise HTTPException(status_code=404, detail="Recommendation not found")
        if decision == "modified" and not modified_text:
            raise HTTPException(status_code=422, detail="Modified decisions require replacement text")
        if recommendation["scientificImpact"] and decision == "accepted":
            raise HTTPException(
                status_code=409,
                detail="Scientific-impact recommendations require an explicit modified text",
            )
        with closing(sqlite3.connect(self.database_path)) as connection, connection:
            connection.execute(
                "INSERT INTO recommendation_decisions VALUES (?, ?, ?, ?, ?, ?)",
                (str(uuid.uuid4()), recommendation_id, principal.workspace_id, decision, modified_text, now_iso()),
            )
        return self.get(principal, analysis_id)

    def artifact_key(self, principal: Principal, analysis_id: str, kind: str) -> str:
        self.get(principal, analysis_id)
        with closing(sqlite3.connect(self.database_path)) as connection:
            row = connection.execute(
                "SELECT object_key FROM analysis_artifacts WHERE analysis_id = ? AND workspace_id = ? AND kind = ?",
                (analysis_id, principal.workspace_id, kind),
            ).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Artifact not found")
        return str(row[0])


LATEX_PATTERN = re.compile(r"\$\$(.+?)\$\$|\\\[(.+?)\\\]|\\\((.+?)\\\)|\$(.+?)\$", re.S)


def _latex_parts(value: object) -> tuple[str, list[str]]:
    text = str(value or "")
    expressions: list[str] = []

    def replace(match: re.Match[str]) -> str:
        expression = next((group for group in match.groups() if group is not None), "").strip()
        if expression:
            expressions.append(expression)
            return " [equation rendered below] "
        return ""

    visible = LATEX_PATTERN.sub(replace, text)
    if (
        not expressions
        and "\\" in text
        and any(token in text for token in ("\\frac", "\\left", "\\right", "\\langle", "\\partial", "\\sum"))
    ):
        expressions.append(text.strip())
        visible = "[equation rendered below]"
    return " ".join(visible.split()), expressions


def _math_png(expression: str, *, color: str = "#1f4e79") -> BytesIO | None:
    value = expression.strip().replace("\n", " ")
    if not value or len(value) > 2_000:
        return None
    output = BytesIO()
    try:
        math_to_image(f"${value}$", output, dpi=180, format="png", color=color)
    except (ValueError, RuntimeError):
        return None
    output.seek(0)
    return output


def _add_review_value(document: Any, label: str, value: object, *, blue: bool = False) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(7)
    heading.paragraph_format.space_after = Pt(2)
    label_run = heading.add_run(label.upper())
    label_run.bold = True
    label_run.font.size = Pt(8)
    label_run.font.color.rgb = RGBColor(95, 102, 105)
    visible, expressions = _latex_parts(value)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(visible or "—")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(31, 78, 121) if blue else RGBColor(0, 0, 0)
    for expression in expressions:
        image = _math_png(expression)
        if image is not None:
            equation = document.add_paragraph()
            equation.alignment = WD_ALIGN_PARAGRAPH.CENTER
            equation.add_run().add_picture(image, width=Inches(5.8))


def _add_suggestion_page(document: Any, title: str, items: list[dict[str, object]]) -> None:
    title_paragraph = document.add_paragraph()
    title_run = title_paragraph.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(11, 63, 92)
    notice = document.add_paragraph(
        "The preceding page is the unchanged original manuscript. Suggestions on this page are color-coded blue; "
        "scientific changes require author validation."
    )
    notice.runs[0].italic = True
    notice.runs[0].font.color.rgb = RGBColor(95, 102, 105)
    for index, item in enumerate(items, 1):
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(12)
        dimension = str(item.get("reviewDimension") or item.get("category") or "Editorial review")
        run = heading.add_run(f"{index}. {dimension.replace('-', ' ').title()}")
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(31, 78, 121)
        _add_review_value(document, "Current manuscript", item.get("originalText"))
        _add_review_value(
            document,
            "Journal/reference pattern",
            item.get("referencePattern") or item.get("journalExpectation") or item.get("basis"),
        )
        _add_review_value(document, "Why this matters", item.get("rationale"))
        _add_review_value(
            document,
            "Author action",
            item.get("authorAction") or item.get("proposedText") or item.get("rationale"),
            blue=True,
        )
        if item.get("proposedText"):
            _add_review_value(document, "Suggested wording", item.get("proposedText"), blue=True)
        if item.get("scientificImpact") or item.get("authorValidationRequired"):
            warning = document.add_paragraph("Author validation required before this scientific change is adopted.")
            warning.runs[0].bold = True
            warning.runs[0].font.color.rgb = RGBColor(176, 31, 31)


def create_pdf_visual_review_docx(original_pdf: bytes, recommendations: list[dict[str, object]]) -> bytes:
    """Preserve a PDF manuscript as source-page images and interleave readable blue review pages."""
    try:
        reader = PdfReader(BytesIO(original_pdf), strict=False)
    except PdfReadError as error:
        raise HTTPException(status_code=422, detail="Original PDF cannot be rendered safely") from error
    if not reader.pages:
        raise HTTPException(status_code=422, detail="Original PDF contains no pages")
    visible = [item for item in recommendations if item.get("decision") != "rejected"]
    by_page: dict[int, list[dict[str, object]]] = {}
    unanchored: list[dict[str, object]] = []
    for item in visible:
        match = re.fullmatch(r"page:(\d+)", str(item.get("anchor", "")))
        if match and 1 <= int(match.group(1)) <= len(reader.pages):
            by_page.setdefault(int(match.group(1)), []).append(item)
        else:
            unanchored.append(item)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.32)
    section.bottom_margin = Inches(0.32)
    section.left_margin = Inches(0.35)
    section.right_margin = Inches(0.35)
    with tempfile.TemporaryDirectory(prefix="article-fit-pdf-") as folder:
        source_path = Path(folder) / "source.pdf"
        source_path.write_bytes(original_pdf)
        prefix = Path(folder) / "page"
        try:
            subprocess.run(
                ["pdftoppm", "-jpeg", "-r", "150", "-jpegopt", "quality=88", str(source_path), str(prefix)],
                check=True,
                capture_output=True,
                timeout=120,
            )
        except (OSError, subprocess.SubprocessError) as error:
            raise HTTPException(status_code=500, detail="PDF source pages could not be rendered") from error
        pages = sorted(Path(folder).glob("page-*.jpg"))
        if len(pages) != len(reader.pages):
            raise HTTPException(status_code=500, detail="PDF source-page rendering was incomplete")
        for page_number, image_path in enumerate(pages, 1):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.add_run().add_picture(str(image_path), width=Inches(7.75))
            document.add_page_break()  # type: ignore[no-untyped-call]
            if by_page.get(page_number):
                _add_suggestion_page(
                    document,
                    f"Article Fit suggestions for source page {page_number}",
                    by_page[page_number],
                )
                document.add_page_break()  # type: ignore[no-untyped-call]
        if unanchored:
            _add_suggestion_page(document, "Article Fit manuscript-level suggestions", unanchored)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def create_docx(text: str, recommendations: list[dict[str, object]], reconstructed: bool) -> bytes:
    paragraphs = [part.strip() for part in text.splitlines() if part.strip()] or ["No extractable manuscript text."]
    section_names = {"abstract", "introduction", "methods", "results", "discussion", "conclusion", "references"}
    body_parts: list[str] = []
    for index, part in enumerate(paragraphs):
        style = "Title" if index == 0 else "Heading1" if part.casefold().rstrip(":") in section_names else "Normal"
        body_parts.append(_word_paragraph(part, style=style))
    body = "".join(body_parts)
    notice = (
        "PDF-source review copy: the original editable Word template cannot be recovered from a PDF. "
        "Use the revised PDF for exact visual fidelity."
        if reconstructed
        else "Article Fit review copy. Original manuscript text remains black."
    )
    body += _word_paragraph(notice, color="C00000", bold=True)
    visible = [item for item in recommendations if item.get("decision") != "rejected"]
    if visible:
        body += _word_paragraph("Color-coded editorial suggestions", color="2E74B5", bold=True, style="Heading1")
    for item in visible:
        proposed = item.get("modifiedText") or item.get("proposedText") or item.get("rationale")
        body += _word_paragraph(
            f"ARTICLE FIT SUGGESTION [{item['category']}] {item['anchor']}: {proposed}",
            color=CATEGORY_COLORS.get(str(item["category"]), "1F4E79"),
        )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:body>{body}<w:sectPr><w:pgSz w:w="12240" w:h="15840"/>'
        '<w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440"/></w:sectPr></w:body></w:document>'
    )
    styles = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/>'
        '<w:pPr><w:spacing w:after="120" w:line="264" w:lineRule="auto"/></w:pPr>'
        '<w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:sz w:val="22"/></w:rPr></w:style>'
        '<w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:basedOn w:val="Normal"/>'
        '<w:pPr><w:spacing w:after="240"/></w:pPr><w:rPr><w:b/><w:color w:val="0B2545"/>'
        '<w:sz w:val="40"/></w:rPr></w:style>'
        '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/>'
        '<w:basedOn w:val="Normal"/><w:pPr><w:keepNext/><w:spacing w:before="320" w:after="160"/></w:pPr>'
        '<w:rPr><w:b/><w:color w:val="2E74B5"/><w:sz w:val="32"/></w:rPr></w:style></w:styles>'
    )
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            '<Override PartName="/word/styles.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/></Types>',
        )
        archive.writestr(
            "_rels/.rels",
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            'Target="word/document.xml"/></Relationships>',
        )
        archive.writestr("word/document.xml", document)
        archive.writestr("word/styles.xml", styles)
        archive.writestr(
            "word/_rels/document.xml.rels",
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>',
        )
    return output.getvalue()


def _word_suggestion_block(item: dict[str, object]) -> str:
    category = str(item.get("reviewDimension") or item.get("category") or "editorial review")
    color = CATEGORY_COLORS.get(str(item.get("category")), "1F4E79")
    original, _ = _latex_parts(item.get("originalText"))
    pattern, _ = _latex_parts(item.get("referencePattern") or item.get("journalExpectation") or item.get("basis"))
    rationale, _ = _latex_parts(item.get("rationale"))
    action, _ = _latex_parts(item.get("authorAction") or item.get("proposedText") or item.get("rationale"))
    proposed, _ = _latex_parts(item.get("modifiedText") or item.get("proposedText"))
    parts = [
        _word_paragraph(f"ARTICLE FIT — {category.replace('-', ' ').upper()}", color=color, bold=True),
    ]
    if original:
        parts.append(_word_paragraph(f"CURRENT MANUSCRIPT: {original}", color="000000"))
    if pattern:
        parts.append(_word_paragraph(f"JOURNAL/REFERENCE PATTERN: {pattern}", color="666666"))
    if rationale:
        parts.append(_word_paragraph(f"WHY THIS MATTERS: {rationale}", color="000000"))
    parts.append(_word_paragraph(f"AUTHOR ACTION: {action}", color=color, bold=True))
    if proposed:
        parts.append(_word_paragraph(f"SUGGESTED WORDING: {proposed}", color=color))
    if item.get("scientificImpact") or item.get("authorValidationRequired"):
        parts.append(
            _word_paragraph(
                "AUTHOR VALIDATION REQUIRED BEFORE ADOPTING THIS SCIENTIFIC CHANGE.",
                color="C00000",
                bold=True,
            )
        )
    return "".join(parts)


def annotate_docx(original: bytes, recommendations: list[dict[str, object]]) -> bytes:
    """Preserve an uploaded DOCX package and add color-coded suggestions near their anchors."""
    try:
        with zipfile.ZipFile(BytesIO(original)) as source:
            document = source.read("word/document.xml").decode("utf-8")
            entries = {name: source.read(name) for name in source.namelist()}
    except (zipfile.BadZipFile, KeyError, UnicodeDecodeError) as error:
        raise HTTPException(status_code=422, detail="Original DOCX cannot be safely annotated") from error

    visible = [item for item in recommendations if item.get("decision") != "rejected"]
    anchored: dict[int, list[dict[str, object]]] = {}
    unanchored: list[dict[str, object]] = []
    for item in visible:
        match = re.fullmatch(r"paragraph:(\d+)", str(item.get("anchor", "")))
        if match:
            anchored.setdefault(int(match.group(1)), []).append(item)
        else:
            unanchored.append(item)

    paragraph_pattern = re.compile(r"<w:p(?:\s[^>]*)?>.*?</w:p>", re.S)
    rebuilt: list[str] = []
    cursor = 0
    logical_index = 0
    for match in paragraph_pattern.finditer(document):
        rebuilt.append(document[cursor : match.end()])
        cursor = match.end()
        if re.search(r"<w:t(?:\s[^>]*)?>.*?</w:t>", match.group(0), re.S):
            logical_index += 1
            for item in anchored.get(logical_index, []):
                rebuilt.append(_word_suggestion_block(item))
    rebuilt.append(document[cursor:])
    document = "".join(rebuilt)

    notes = ""
    if unanchored:
        notes += _word_paragraph("Article Fit — manuscript-level suggestions", color="2E74B5", bold=True)
    for item in unanchored:
        notes += _word_suggestion_block(item)
    insertion = document.rfind("<w:sectPr")
    if insertion < 0:
        insertion = document.rfind("</w:body>")
    if insertion < 0:
        raise HTTPException(status_code=422, detail="Original DOCX body cannot be located")
    entries["word/document.xml"] = (document[:insertion] + notes + document[insertion:]).encode()
    core = entries.get("docProps/core.xml")
    if core is not None:
        core = re.sub(rb"(<dc:creator[^>]*>).*?(</dc:creator>)", rb"\1\2", core, flags=re.S)
        core = re.sub(rb"(<cp:lastModifiedBy[^>]*>).*?(</cp:lastModifiedBy>)", rb"\1\2", core, flags=re.S)
        entries["docProps/core.xml"] = core
    entries.pop("docProps/custom.xml", None)
    if "_rels/.rels" in entries:
        entries["_rels/.rels"] = re.sub(
            rb"<Relationship\b[^>]*Target=\"docProps/custom\.xml\"[^>]*/>", b"", entries["_rels/.rels"]
        )
    if "[Content_Types].xml" in entries:
        entries["[Content_Types].xml"] = re.sub(
            rb"<Override\b[^>]*PartName=\"/docProps/custom\.xml\"[^>]*/>",
            b"",
            entries["[Content_Types].xml"],
        )
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name, content in entries.items():
            target.writestr(name, content)
    annotated = output.getvalue()
    rendered_equations: list[tuple[str, str]] = []
    for item in visible:
        proposed = item.get("modifiedText") or item.get("proposedText") or item.get("rationale")
        _, expressions = _latex_parts(proposed)
        rendered_equations.extend((str(item.get("anchor", "document")), expression) for expression in expressions)
    if not rendered_equations:
        return annotated
    document_with_math = Document(BytesIO(annotated))
    heading = document_with_math.add_paragraph()
    run = heading.add_run("Article Fit — rendered equations in suggested revisions")
    run.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)
    for anchor, expression in rendered_equations:
        image = _math_png(expression)
        if image is None:
            continue
        paragraph = document_with_math.add_paragraph(anchor)
        paragraph.runs[0].font.color.rgb = RGBColor(95, 102, 105)
        equation = document_with_math.add_paragraph()
        equation.alignment = WD_ALIGN_PARAGRAPH.CENTER
        equation.add_run().add_picture(image, width=Inches(5.8))
    rendered = BytesIO()
    document_with_math.save(rendered)
    return rendered.getvalue()


def validate_invariant_preservation(original: str, revised: str) -> None:
    """Stop artifact generation if an original scientific token disappears."""
    original_invariants = scientific_invariants(original)
    revised_invariants = scientific_invariants(revised)
    for group, values in original_invariants.items():
        revised_values = revised_invariants[group]
        if any(value not in revised_values for value in values):
            raise HTTPException(status_code=409, detail=f"Scientific invariant validation failed: {group}")


def validate_proposal_parity(docx: bytes, revised_pdf: bytes, recommendations: list[dict[str, object]]) -> None:
    """Require each visible proposal to be represented in both revision formats."""
    with zipfile.ZipFile(BytesIO(docx)) as archive:
        docx_xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
    try:
        root = ET.fromstring(docx_xml)
        docx_text = " ".join(node.text or "" for node in root.iter() if node.tag.endswith("}t"))
    except ET.ParseError:
        docx_text = unescape(re.sub(r"<[^>]+>", " ", docx_xml))
    try:
        pdf_text = " ".join(page.extract_text() or "" for page in PdfReader(BytesIO(revised_pdf), strict=False).pages)
    except (PdfReadError, ValueError, TypeError) as error:
        raise HTTPException(status_code=500, detail="Revised PDF cannot be read") from error
    normalized_docx = " ".join(docx_text.split())
    normalized_pdf = " ".join(pdf_text.split())
    for item in recommendations:
        if item.get("decision") == "rejected":
            continue
        candidates = [
            str(value)
            for value in (
                item.get("authorAction"),
                item.get("modifiedText"),
                item.get("proposedText"),
                item.get("rationale"),
            )
            if value
        ]
        probes = [" ".join(candidate.split())[:80] for candidate in candidates]
        if not any(probe in normalized_docx and probe in normalized_pdf for probe in probes):
            raise HTTPException(status_code=500, detail="DOCX/PDF recommendation parity validation failed")


def _word_paragraph(text: str, color: str = "000000", bold: bool = False, style: str = "Normal") -> str:
    text = _sanitize_xml_text(text)
    bold_xml = "<w:b/>" if bold else ""
    return (
        f'<w:p><w:pPr><w:pStyle w:val="{style}"/></w:pPr><w:r><w:rPr>'
        f'<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:color w:val="{color}"/>{bold_xml}'
        f'</w:rPr><w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p>'
    )


def _sanitize_xml_text(value: object) -> str:
    text = str(value or "")
    return "".join(
        character
        for character in text
        if character in "\t\n\r"
        or 0x20 <= ord(character) <= 0xD7FF
        or 0xE000 <= ord(character) <= 0xFFFD
        or 0x10000 <= ord(character) <= 0x10FFFF
    )


def create_pdf(title: str, lines: Iterable[str]) -> bytes:
    wrapped = [piece for line in lines for piece in (textwrap.wrap(line, width=90) or [""])]
    pages = [wrapped[index : index + 48] for index in range(0, len(wrapped), 48)] or [[]]
    objects: list[bytes] = []
    page_ids: list[int] = []
    font_id = 3
    objects.extend([b"", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"])
    for page_lines in pages:
        content_parts = ["BT /F1 16 Tf 72 740 Td", f"({_pdf_escape(title)}) Tj", "0 -26 Td /F1 10 Tf"]
        for line in page_lines:
            content_parts.extend([f"({_pdf_escape(line)}) Tj", "0 -14 Td"])
        content_parts.append("ET")
        stream = "\n".join(content_parts).encode("latin-1", errors="replace")
        content_id = len(objects) + 1
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
        page_id = len(objects) + 1
        page_ids.append(page_id)
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>"
            ).encode()
        )
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return bytes(output)


def _pdf_escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def validate_artifacts(docx: bytes, revised_pdf: bytes, report_pdf: bytes) -> dict[str, object]:
    try:
        with zipfile.ZipFile(BytesIO(docx)) as archive:
            document_xml = archive.read("word/document.xml")
            ET.fromstring(document_xml)
    except (zipfile.BadZipFile, KeyError, ET.ParseError) as error:
        raise HTTPException(status_code=500, detail="Generated DOCX failed structural validation") from error
    try:
        revised_reader = PdfReader(BytesIO(revised_pdf), strict=False)
        report_reader = PdfReader(BytesIO(report_pdf), strict=False)
        if not revised_reader.pages or not report_reader.pages:
            raise ValueError("empty PDF")
    except (PdfReadError, ValueError, TypeError):
        raise HTTPException(status_code=500, detail="Generated PDF failed structural validation") from None
    forbidden = (b"Bearer ", b"local-invite-token", b"/tmp/journal-matcher")
    combined = document_xml + revised_pdf + report_pdf
    if any(secret in combined for secret in forbidden):
        raise HTTPException(status_code=500, detail="Artifact privacy validation failed")
    return {"structural": True, "privacy": True, "sourceStructurePreserved": True}


def store_artifact_set(
    store: FoundationStore,
    repository: AnalysisRepository,
    principal: Principal,
    analysis_id: str,
    artifacts: dict[str, bytes],
    validation: dict[str, object],
) -> dict[str, object]:
    with closing(sqlite3.connect(repository.database_path)) as connection, connection:
        for kind, content in artifacts.items():
            object_key = f"{principal.workspace_id}/artifacts/{analysis_id}/{kind}"
            path = store.object_root / object_key
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(content)
            path.chmod(0o600)
            connection.execute(
                "INSERT OR REPLACE INTO analysis_artifacts VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    stable_id(analysis_id, kind),
                    analysis_id,
                    principal.workspace_id,
                    kind,
                    object_key,
                    f"sha256:{hashlib.sha256(content).hexdigest()}",
                    json.dumps(validation, sort_keys=True),
                    now_iso(),
                ),
            )
    return repository.get(principal, analysis_id)
